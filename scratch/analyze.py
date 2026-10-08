import docx
import sys

def analyze_docx(filepath):
    print(f"--- Analyzing {filepath} ---")
    doc = docx.Document(filepath)
    for i, p in enumerate(doc.paragraphs[:20]):
        print(f"Para {i}: {p.text.strip()[:100]}")
    
    print("Tables:")
    for i, t in enumerate(doc.tables[:2]):
        print(f"Table {i} has {len(t.rows)} rows and {len(t.columns)} cols")
        for r_idx, row in enumerate(t.rows[:2]):
            cells = [c.text.strip().replace('\n', ' ')[:50] for c in row.cells]
            print(f"  Row {r_idx}: {cells}")
    print("\n")

analyze_docx(r"C:\Users\USUARIO\Desktop\AI\Anteproyecto_APA_Interactivo_APA7.docx")
analyze_docx(r"C:\Users\USUARIO\Desktop\AI\ANTEPROYECTO_V18_FINAL (1).docx")
analyze_docx(r"C:\Users\USUARIO\Desktop\AI\Metodologia Web APA DLM.docx")
